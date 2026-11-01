import io
import json
import zipfile
from types import SimpleNamespace
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.conftest import ADMIN_PW, TERMS, login, make_user, oracle_export, upload


def converted(
    c: TestClient, name: str = "Export_20240101_1200.zip", data: bytes | None = None
) -> dict[str, Any]:
    ws = upload(c, name, data or oracle_export())
    r = c.post(f"/api/workspaces/{ws['id']}/convert", json={})
    assert r.status_code == 202, r.text
    job = c.get(f"/api/jobs/{r.json()['id']}").json()
    assert job["status"] == "succeeded", job
    detail = c.get(f"/api/workspaces/{ws['id']}").json()
    assert detail["status"] == "converted"
    return detail  # type: ignore[no-any-return]


def ds_by_name(ws: dict[str, Any], name: str) -> dict[str, Any]:
    return next(d for d in ws["datasets"] if d["name"] == name)


def sse_events(body: str) -> list[tuple[str, Any]]:
    out = []
    for frame in body.strip().split("\n\n"):
        lines = dict(line.split(": ", 1) for line in frame.splitlines())
        out.append((lines["event"], json.loads(lines["data"])))
    return out


# --- auth & security ----------------------------------------------------------------------


def test_requires_login(client: TestClient) -> None:
    assert client.get("/api/workspaces").status_code == 401
    assert client.get("/api/auth/me").status_code == 401


def test_login_logout_and_cookie_flags(client: TestClient) -> None:
    r = client.post("/api/auth/login", json={"username": "admin", "password": ADMIN_PW})
    cookie = r.headers["set-cookie"]
    assert "HttpOnly" in cookie and "SameSite=strict" in cookie and "Path=/api" in cookie
    assert client.get("/api/auth/me").json()["role"] == "admin"
    client.post("/api/auth/logout")
    assert client.get("/api/auth/me").status_code == 401


def test_lockout_after_repeated_failures(admin: TestClient) -> None:
    make_user(admin, "carol")
    admin.post("/api/auth/logout")
    for _ in range(5):
        r = admin.post("/api/auth/login", json={"username": "carol", "password": "nope"})
        assert r.status_code == 401
    r = admin.post("/api/auth/login", json={"username": "carol", "password": "User-Pass-2026!"})
    assert r.status_code == 423
    login(admin, "admin", ADMIN_PW)
    carol = next(u for u in admin.get("/api/admin/users").json() if u["username"] == "carol")
    admin.post(f"/api/admin/users/{carol['id']}/unlock")
    admin.post("/api/auth/logout")
    login(admin, "carol", "User-Pass-2026!")


def test_unknown_user_and_wrong_password_look_identical(client: TestClient) -> None:
    a = client.post("/api/auth/login", json={"username": "ghost", "password": "x"})
    b = client.post("/api/auth/login", json={"username": "admin", "password": "x"})
    assert (a.status_code, a.json()) == (b.status_code, b.json())


def test_cross_origin_writes_blocked(admin: TestClient) -> None:
    r = admin.post("/api/workspaces/x/convert", json={}, headers={"Origin": "https://evil.test"})
    assert r.status_code == 403
    r = admin.post("/api/auth/logout", headers={"Sec-Fetch-Site": "cross-site"})
    assert r.status_code == 403
    ok = admin.post("/api/auth/logout", headers={"Origin": "http://localhost:5173"})
    assert ok.status_code == 200


def test_security_headers(client: TestClient) -> None:
    r = client.get("/api/health")
    assert r.headers["x-content-type-options"] == "nosniff"
    assert "frame-ancestors 'none'" in r.headers["content-security-policy"]
    assert r.headers["cache-control"] == "no-store"
    bad = client.get("/api/health", headers={"X-Request-ID": "evil\nlog"})
    assert "\n" not in bad.headers["x-request-id"]


def test_admin_only_endpoints(admin: TestClient) -> None:
    make_user(admin, "dave")
    admin.post("/api/auth/logout")
    login(admin, "dave", "User-Pass-2026!")
    assert admin.get("/api/admin/users").status_code == 403
    assert (
        admin.post("/api/admin/users", json={"username": "x", "password": "y" * 12}).status_code
        == 403
    )


def test_password_policy_and_reset_flow(admin: TestClient) -> None:
    r = admin.post("/api/admin/users", json={"username": "erin", "password": "short"})
    assert r.status_code == 422
    r = admin.post("/api/admin/users", json={"username": "erin", "password": "erin-password-1"})
    assert r.status_code == 422 and "username" in r.json()["detail"]
    make_user(admin, "erin")
    erin = next(u for u in admin.get("/api/admin/users").json() if u["username"] == "erin")

    admin.post("/api/auth/logout")
    assert admin.post("/api/auth/reset-requests", json={"username": "erin"}).status_code == 202
    assert admin.post("/api/auth/reset-requests", json={"username": "nobody"}).status_code == 202

    login(admin, "admin", ADMIN_PW)
    assert [r["username"] for r in admin.get("/api/admin/reset-requests").json()] == [
        "nobody",
        "erin",
    ]
    token = admin.post(f"/api/admin/users/{erin['id']}/reset-token").json()["token"]
    admin.post("/api/auth/logout")
    r = admin.post("/api/auth/reset", json={"token": token, "new_password": "Brand-New-Pass-1"})
    assert r.status_code == 200
    reuse = admin.post("/api/auth/reset", json={"token": token, "new_password": "Another-Pass-22"})
    assert reuse.status_code == 400
    login(admin, "erin", "Brand-New-Pass-1")


def test_admin_cannot_modify_self(admin: TestClient) -> None:
    me = admin.get("/api/auth/me").json()
    assert admin.patch(f"/api/admin/users/{me['id']}", json={"role": "user"}).status_code == 400


# --- workspace flow -----------------------------------------------------------------------


def test_upload_scan_groups_and_notes(admin: TestClient) -> None:
    ws = upload(admin, "Export_20240101_1200.zip", oracle_export())
    groups = {g["code"]: g for g in ws["groups"]}
    assert set(groups) == {"AR", "FUN", "GL", "INV"}
    assert groups["AR"]["module"] == "Receivables"
    assert groups["FUN"]["domain"]
    assert ws["status"] == "scanned" and len(ws["sources"]) == 5


def test_upload_rejections(admin: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    r = admin.post("/api/workspaces", params={"filename": "x.exe"}, content=b"MZ")
    assert r.status_code == 422
    r = admin.post("/api/workspaces", params={"filename": "x.zip"}, content=b"")
    assert r.status_code == 422
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "max_upload_mb", 0)
    r = admin.post("/api/workspaces", params={"filename": "x.xml"}, content=b"<a/>")
    assert r.status_code == 413
    ws_root = get_settings().workspaces_dir
    assert list(ws_root.iterdir()) == []  # failed uploads leave nothing behind


def test_convert_preview_edit_export(admin: TestClient) -> None:
    ws = converted(admin)
    names = sorted(d["name"] for d in ws["datasets"])
    assert names == [
        "AR_PAYMENT_TERMS.CSV",
        "FUN_BUSINESS_UNIT.CSV",
        "INVENTORYORGANIZATIONCHILD.CSV",
        "INVENTORYORGANIZATIONMASTER.CSV",
        "LEDGERS.CSV",
    ]
    terms = ds_by_name(ws, "AR_PAYMENT_TERMS.CSV")
    assert (terms["rows"], terms["columns"], terms["group"]) == (3, 4, "AR")

    page = admin.get(f"/api/datasets/{terms['id']}", params={"limit": 2}).json()
    assert [c["name"] for c in page["columns"]] == ["TermId", "Name", "DueDays", "CreatedBy"]
    assert page["total"] == 3 and len(page["rows"]) == 2
    found = admin.get(f"/api/datasets/{terms['id']}", params={"q": "net6"}).json()
    assert found["total"] == 1 and found["rows"][0][1] == "NET60"

    r = admin.patch(
        f"/api/datasets/{terms['id']}",
        json={
            "cells": [{"row": 0, "column": "Name", "value": '=HYPERLINK("http://x","click")'}],
            "append_rows": [{"TermId": "9", "Name": "NET90"}],
            "delete_rows": [2],
        },
    )
    assert r.status_code == 200 and r.json()["rows"] == 3 and r.json()["edited"]
    bad = admin.patch(
        f"/api/datasets/{terms['id']}",
        json={"cells": [{"row": 99, "column": "Name", "value": "x"}]},
    )
    assert bad.status_code == 422

    x = admin.get(f"/api/datasets/{terms['id']}/download", params={"format": "xlsx"})
    cell = load_workbook(io.BytesIO(x.content)).active["B2"]  # type: ignore[index]
    assert cell.value.startswith("=HYPERLINK") and cell.data_type == "s"  # text, not formula

    r = admin.post(
        f"/api/workspaces/{ws['id']}/export", json={"format": "csv", "preserve_structure": True}
    )
    assert r.status_code == 200
    members = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert "Export_20240101_1200.zip/AR_PAYMENT_TERM.zip/AR_PAYMENT_TERMS.CSV" in members


def test_exclusions_and_group_filter(admin: TestClient) -> None:
    ws = upload(admin, "Export_20240101_1200.zip", oracle_export())
    gl = [s["id"] for s in ws["sources"] if s["group"] == "GL"]
    admin.post(f"/api/workspaces/{ws['id']}/exclusions", json={"source_ids": gl, "excluded": True})
    admin.post(f"/api/workspaces/{ws['id']}/convert", json={"groups": ["ar", "GL"]})
    detail = admin.get(f"/api/workspaces/{ws['id']}").json()
    assert [d["name"] for d in detail["datasets"]] == ["AR_PAYMENT_TERMS.CSV"]


def test_business_structure(admin: TestClient) -> None:
    ws = converted(admin)
    s = admin.get(f"/api/workspaces/{ws['id']}/structure").json()
    assert s["status"] == "ready"
    ledgers = {n["name"]: n for n in s["root"]["children"]}
    assert set(ledgers) == {"UK Ledger", "US Ledger"}
    us_ou = ledgers["US Ledger"]["children"][0]["children"][0]
    assert us_ou["name"] == "US Ops" and [o["name"] for o in us_ou["children"]] == ["Dallas DC"]
    assert s["counts"]["inventory_org"] == 2  # Leeds de-duplicated, Orphan pruned
    assert any("Skipped 1" in d for d in s["diagnostics"])


def test_comparison_flow(admin: TestClient) -> None:
    left = converted(admin)
    changed = [
        dict(TERMS[0], CreatedBy="zed"),
        dict(TERMS[1], DueDays="45"),
        {"TermId": "4", "Name": "NET90", "DueDays": "90", "CreatedBy": "x"},
    ]
    right = converted(admin, "Export_20240202_0900.zip", oracle_export(changed))
    r = admin.post(
        "/api/comparisons",
        json={"left_workspace_id": left["id"], "right_workspace_id": right["id"]},
    )
    assert r.status_code == 202
    assert admin.get(f"/api/jobs/{r.json()['id']}").json()["status"] == "succeeded"
    cmp_id = r.json()["subject_id"]
    detail = admin.get(f"/api/comparisons/{cmp_id}").json()
    assert detail["summary"] == {"SAME": 4, "MODIFIED": 1, "ADDED": 0, "REMOVED": 0}
    item = next(i for i in detail["items"] if i["status"] == "MODIFIED")
    assert (item["added"], item["removed"], item["modified"]) == (1, 1, 1)
    diff = admin.get(f"/api/comparisons/{cmp_id}/items/{item['id']}/diff").json()
    assert diff["stats"] == {"equal": 1, "modified": 1, "removed": 1, "added": 1}
    assert diff["header"][diff["ignored"][0]] == "CreatedBy"


def test_compare_requires_converted(admin: TestClient) -> None:
    a = upload(admin, "A_X.xml", b"<R><ROW><a>1</a></ROW></R>")
    b = converted(admin)
    r = admin.post(
        "/api/comparisons", json={"left_workspace_id": a["id"], "right_workspace_id": b["id"]}
    )
    assert r.status_code == 409


def test_other_users_get_404_everywhere(admin: TestClient) -> None:
    ws = converted(admin)
    ds = ws["datasets"][0]
    other = converted(admin, "Export_20240202_0900.zip")
    cmp = admin.post(
        "/api/comparisons", json={"left_workspace_id": ws["id"], "right_workspace_id": other["id"]}
    ).json()
    make_user(admin, "mallory")
    admin.post("/api/auth/logout")
    login(admin, "mallory", "User-Pass-2026!")
    for method, url in [
        ("GET", f"/api/workspaces/{ws['id']}"),
        ("DELETE", f"/api/workspaces/{ws['id']}"),
        ("GET", f"/api/datasets/{ds['id']}"),
        ("GET", f"/api/datasets/{ds['id']}/download"),
        ("DELETE", f"/api/datasets/{ds['id']}"),
        ("GET", f"/api/workspaces/{ws['id']}/structure"),
        ("GET", f"/api/workspaces/{ws['id']}/ai/messages"),
        ("GET", f"/api/jobs/{cmp['id']}"),
        ("GET", f"/api/comparisons/{cmp['subject_id']}"),
    ]:
        assert admin.request(method, url).status_code == 404, url
    assert admin.patch(f"/api/datasets/{ds['id']}", json={}).status_code == 404
    assert admin.get("/api/workspaces").json() == []
    mine = upload(admin, "AP_X.xml", b"<R><ROW><a>1</a></ROW></R>")
    r = admin.post(
        "/api/comparisons", json={"left_workspace_id": mine["id"], "right_workspace_id": ws["id"]}
    )
    assert r.status_code == 404


def test_delete_workspace_removes_files(admin: TestClient) -> None:
    from app.config import get_settings

    ws = converted(admin)
    assert (get_settings().workspaces_dir / ws["id"]).exists()
    assert admin.delete(f"/api/workspaces/{ws['id']}").status_code == 200
    assert not (get_settings().workspaces_dir / ws["id"]).exists()
    assert admin.get(f"/api/workspaces/{ws['id']}").status_code == 404


# --- AI ------------------------------------------------------------------------------------


def test_chat_retrieval_mode_without_llm(admin: TestClient) -> None:
    ws = converted(admin)
    assert (
        admin.post(f"/api/workspaces/{ws['id']}/ai/chat", json={"question": "x"}).status_code == 409
    )
    job = admin.post(f"/api/workspaces/{ws['id']}/ai/index").json()
    assert admin.get(f"/api/jobs/{job['id']}").json()["status"] == "succeeded"
    status_ = admin.get(f"/api/workspaces/{ws['id']}/ai").json()
    assert status_["indexed_rows"] == 3 + 3 + 2 + 2 + 2 and not status_["llm_configured"]

    r = admin.post(
        f"/api/workspaces/{ws['id']}/ai/chat",
        json={"question": "Which inventory org belongs to US Ops?"},
    )
    events = sse_events(r.text)
    assert [e for e, _ in events] == ["citations", "token", "done"]
    cites = events[0][1]
    assert cites and any("Dallas DC" in c["snippet"] for c in cites)
    assert events[2][1]["mode"] == "retrieval"
    history = admin.get(f"/api/workspaces/{ws['id']}/ai/messages").json()
    assert [m["role"] for m in history] == ["user", "assistant"]
    md = admin.get(f"/api/workspaces/{ws['id']}/ai/transcript").text
    assert "Which inventory org" in md


def test_fts_query_is_injection_safe(admin: TestClient) -> None:
    from app.services.search import fts_query

    assert fts_query('NEAR(") OR workspace : * AND "') == '"near" OR "workspace"'
    assert fts_query("the of and") is None
    ws = converted(admin)
    admin.post(f"/api/workspaces/{ws['id']}/ai/index")
    r = admin.post(
        f"/api/workspaces/{ws['id']}/ai/chat", json={"question": 'workspace : * OR "" NEAR( ('}
    )
    assert r.status_code == 200 and sse_events(r.text)[-1][0] == "done"


def test_chat_with_llm_streams_and_grounds(
    admin: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.config import get_settings
    from app.services import assistant

    ws = converted(admin)
    admin.post(f"/api/workspaces/{ws['id']}/ai/index")
    s = get_settings()
    monkeypatch.setattr(s, "azure_openai_endpoint", "https://example.openai.azure.com")
    monkeypatch.setattr(s, "azure_openai_api_key", "k")
    monkeypatch.setattr(s, "azure_openai_chat_model", "gpt-test")
    seen: dict[str, Any] = {}

    def chunk(text: str | None) -> SimpleNamespace:
        return SimpleNamespace(
            usage=None, choices=[SimpleNamespace(delta=SimpleNamespace(content=text))]
        )

    class FakeCompletions:
        def create(self, **kw: Any) -> list[SimpleNamespace]:
            seen.update(kw)
            return [chunk("NET60 is due in "), chunk("**60** days [1]."), chunk(None)]

    fake = SimpleNamespace(chat=SimpleNamespace(completions=FakeCompletions()))
    monkeypatch.setattr(assistant, "_client", lambda: fake)
    r = admin.post(f"/api/workspaces/{ws['id']}/ai/chat", json={"question": "NET60 due days"})
    events = sse_events(r.text)
    assert "".join(d for e, d in events if e == "token") == "NET60 is due in **60** days [1]."
    assert events[-1][1]["mode"] == "llm"
    assert seen["model"] == "gpt-test" and seen["stream"] is True
    system, user = seen["messages"][0]["content"], seen["messages"][-1]["content"]
    assert "untrusted" in system and user.startswith("DATA:\n[1]") and "NET60" in user


def test_llm_failure_is_reported(admin: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.config import get_settings
    from app.services import assistant

    ws = converted(admin)
    admin.post(f"/api/workspaces/{ws['id']}/ai/index")
    s = get_settings()
    monkeypatch.setattr(s, "azure_openai_endpoint", "https://e")
    monkeypatch.setattr(s, "azure_openai_api_key", "k")
    monkeypatch.setattr(s, "azure_openai_chat_model", "m")

    class Boom:
        def create(self, **kw: Any) -> None:
            raise TimeoutError("slow")

    monkeypatch.setattr(
        assistant, "_client", lambda: SimpleNamespace(chat=SimpleNamespace(completions=Boom()))
    )
    r = admin.post(f"/api/workspaces/{ws['id']}/ai/chat", json={"question": "NET60"})
    assert sse_events(r.text)[-1] == ("error", {"message": "AI service error: TimeoutError"})


def test_edit_reindexes_search(admin: TestClient) -> None:
    ws = converted(admin)
    admin.post(f"/api/workspaces/{ws['id']}/ai/index")
    terms = ds_by_name(ws, "AR_PAYMENT_TERMS.CSV")
    admin.patch(
        f"/api/datasets/{terms['id']}",
        json={"cells": [{"row": 2, "column": "Name", "value": "ZEBRATERM"}]},
    )
    r = admin.post(f"/api/workspaces/{ws['id']}/ai/chat", json={"question": "zebraterm"})
    assert "ZEBRATERM" in sse_events(r.text)[0][1][0]["snippet"]


def test_interrupted_jobs_recovered(admin: TestClient) -> None:
    from app.db import new_session
    from app.models import Job, Workspace
    from app.services import jobs

    ws = upload(admin, "AR_X.xml", b"<R><ROW><a>1</a></ROW></R>")
    with new_session() as db:
        w = db.get(Workspace, ws["id"])
        assert w is not None
        w.status = "converting"
        db.add(Job(owner_id=1, kind="convert", subject_id=ws["id"], status="running"))
        db.commit()
    assert jobs.recover_interrupted() == 1
    assert admin.get(f"/api/workspaces/{ws['id']}").json()["status"] == "failed"
